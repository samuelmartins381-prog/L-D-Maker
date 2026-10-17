#!/usr/bin/env python3
"""Builds the invented, non-confidential test sources of projects/mouvement-mecanique (ADR 0013).

Three formats on purpose, to exercise the three source readers: a PDF guide (written as DOCX, converted
by LibreOffice), a PPTX sales deck and a DOCX of field notes. General horology facts only: no brand,
no real calibre. Uses python-docx, python-pptx, Pillow and LibreOffice from the environment.

    python3 scripts/demo-sources/make_mouvement_sources.py projects/mouvement-mecanique/sources
"""
import math
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Cm as PCm, Pt as PPt

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
FONT = "/usr/share/fonts/opentype/inter/Inter-SemiBold.otf"

# ---------------------------------------------------------------- schematic (numbers only, no words)


def gear(d, cx, cy, r, teeth, fill, outline="#3b3631"):
    pts = []
    for i in range(teeth * 2):
        a = math.pi * i / teeth
        rr = r if i % 2 == 0 else r * 0.88
        pts.append((cx + rr * math.cos(a), cy + rr * math.sin(a)))
    d.polygon(pts, fill=fill, outline=outline)
    d.ellipse((cx - 6, cy - 6, cx + 6, cy + 6), fill=outline)


def schematic(path: Path):
    """The five functions of a mechanical movement, numbered 1 to 5 (legend lives in the text)."""
    w, h = 1600, 900
    img = Image.new("RGB", (w, h), "#fffdf9")
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, 44)
    # 1 barrel (energy)
    d.ellipse((90, 300, 410, 620), fill="#e8dccd", outline="#3b3631", width=6)
    for k in range(6):  # mainspring
        r = 40 + k * 20
        d.arc((250 - r, 460 - r, 250 + r, 460 + r), 0, 300, fill="#8a4b2a", width=5)
    # 2 gear train (transmission)
    gear(d, 560, 360, 120, 24, "#efe4d6")
    gear(d, 760, 520, 90, 18, "#efe4d6")
    gear(d, 930, 380, 70, 14, "#efe4d6")
    # 3 escapement: escape wheel + pallet fork
    gear(d, 1110, 540, 70, 15, "#f1e4da")
    d.line((1180, 470, 1290, 380), fill="#3b3631", width=14)
    d.line((1180, 470, 1150, 400), fill="#3b3631", width=10)
    d.line((1180, 470, 1230, 520), fill="#3b3631", width=10)
    # 4 balance and hairspring (regulation)
    d.ellipse((1260, 150, 1520, 410), outline="#3b3631", width=14)
    for k in range(5):
        r = 18 + k * 18
        d.arc((1390 - r, 280 - r, 1390 + r, 280 + r), 20, 340, fill="#2f5aa8", width=4)
    # 5 hands (display)
    d.ellipse((470, 640, 650, 820), outline="#3b3631", width=6)
    d.line((560, 730, 560, 660), fill="#3b3631", width=8)
    d.line((560, 730, 615, 760), fill="#3b3631", width=8)
    # flow arrows
    for x1, y1, x2, y2 in [(410, 420, 440, 400), (650, 430, 680, 470), (850, 470, 870, 430),
                           (1000, 400, 1050, 500), (1240, 470, 1280, 400), (560, 480, 560, 630)]:
        d.line((x1, y1, x2, y2), fill="#8a4b2a", width=6)
    # numbered markers
    for n, (x, y) in enumerate([(250, 250), (560, 200), (1110, 650), (1390, 110), (700, 730)], 1):
        d.ellipse((x - 34, y - 34, x + 34, y + 34), fill="#1d1b18")
        d.text((x, y), str(n), fill="#fffdf9", font=font, anchor="mm")
    img.save(path)


# ---------------------------------------------------------------- guide (DOCX → PDF)

GUIDE = [
    ("h1", "Le mouvement mécanique — Guide du conseiller"),
    ("p", "Contenu de démonstration inventé pour tester le générateur de modules. Faits horlogers généraux, sans marque ni calibre réel."),
    ("h2", "1. Une petite machine sans pile"),
    ("p", "Une montre mécanique fonctionne sans pile ni électronique. C'est une machine miniature qui stocke de l'énergie dans un ressort, la libère par petites quantités parfaitement régulières, compte ces battements et les traduit en heures, minutes et secondes."),
    ("p", "Pour l'expliquer simplement, on distingue cinq fonctions, toujours dans le même ordre : l'énergie, la transmission, la distribution, la régulation et l'affichage. Le schéma ci-dessous les numérote de 1 à 5."),
    ("img", "schema-mouvement.png"),
    ("p", "1 : le barillet (énergie). 2 : le rouage (transmission). 3 : l'échappement (distribution). 4 : le balancier-spiral (régulation). 5 : les aiguilles (affichage)."),
    ("h2", "2. L'énergie : le barillet"),
    ("p", "L'énergie est stockée dans un long ressort enroulé, le ressort moteur, logé dans un tambour denté appelé barillet. Remonter la montre, c'est armer ce ressort. En se détendant, il fait tourner le barillet, qui entraîne tout le mouvement."),
    ("p", "Sur une montre à remontage manuel, on arme le ressort en tournant la couronne. Sur une montre automatique, une masse oscillante, le rotor, tourne au gré des mouvements du poignet et remonte le ressort. Un dispositif de glissement évite d'armer le ressort au-delà de sa limite."),
    ("p", "La réserve de marche est la durée pendant laquelle la montre fonctionne une fois le ressort entièrement armé. Pour un mouvement courant, elle est souvent de 38 à 50 heures ; certains mouvements dépassent 70 heures, voire plusieurs jours."),
    ("h2", "3. La transmission : le rouage"),
    ("p", "Le rouage est une suite de roues dentées qui transmet l'énergie du barillet jusqu'à l'échappement. Chaque roue tourne à une vitesse précise : la roue de centre fait un tour par heure et porte l'aiguille des minutes ; la roue de secondes fait un tour par minute."),
    ("p", "Les axes des roues tournent dans des rubis synthétiques. Ces pierres très dures servent de paliers : elles réduisent les frottements et l'usure. Leur nombre dépend surtout de la construction du mouvement et de ses fonctions : un nombre élevé de rubis n'est pas, en soi, un gage de qualité."),
    ("h2", "4. La distribution : l'échappement"),
    ("p", "L'échappement libère l'énergie du rouage par petites impulsions. Le plus répandu est l'échappement à ancre suisse : la roue d'échappement est bloquée puis libérée, dent après dent, par une ancre qui bascule de gauche à droite. À chaque bascule, l'ancre donne une petite impulsion au balancier. C'est ce mécanisme qui produit le tic-tac."),
    ("h2", "5. La régulation : le balancier-spiral"),
    ("p", "Le balancier-spiral est le cœur de la montre. Le balancier est un anneau qui oscille d'un côté puis de l'autre, ramené par un ressort très fin, le spiral. Ses oscillations ont une durée constante : c'est lui qui donne la mesure du temps."),
    ("p", "Sa fréquence s'exprime en alternances par heure ou en hertz. 28 800 alternances par heure correspondent à 4 hertz, soit 8 alternances par seconde ; c'est une fréquence très répandue. On trouve aussi 21 600 alternances par heure (3 hertz) et 36 000 (5 hertz). Une fréquence plus élevée rend la marche plus stable face aux perturbations, mais demande plus d'énergie et use davantage les pièces."),
    ("p", "C'est aussi pourquoi l'aiguille des secondes d'une montre mécanique semble glisser : à 4 hertz, elle avance par 8 petits pas par seconde, alors que celle d'une montre à quartz saute une fois par seconde."),
    ("h2", "6. L'affichage"),
    ("p", "Un petit train de roues, la minuterie, divise par douze la vitesse de l'aiguille des minutes pour entraîner l'aiguille des heures. Les fonctions supplémentaires, comme la date ou le chronographe, sont appelées des complications."),
    ("h2", "7. Précision : ce qu'il faut savoir"),
    ("p", "Une montre mécanique avance ou retarde de quelques secondes par jour. Le titre de chronomètre est attribué après des tests officiels de plusieurs jours : la marche moyenne journalière doit rester entre −4 et +6 secondes. Une montre à quartz courante dérive, elle, d'environ 15 secondes par mois."),
    ("p", "La différence vient de l'oscillateur : un balancier bat quelques fois par seconde, un quartz vibre 32 768 fois par seconde. La marche d'une mécanique varie aussi avec la position de la montre, la température, le magnétisme et l'état d'armage du ressort."),
    ("h2", "8. Les ennemis de la montre"),
    ("p", "Le magnétisme : un spiral aimanté colle légèrement sur lui-même et la montre peut se mettre à avancer fortement, de plusieurs dizaines de secondes par jour. Haut-parleurs, fermoirs magnétiques de sacs ou tablettes peuvent en être la cause. Une démagnétisation chez l'horloger règle le problème en quelques minutes."),
    ("p", "Les chocs : un système antichoc protège les pivots fragiles du balancier. Il amortit les chocs du quotidien, pas une chute sur un sol dur."),
    ("p", "L'eau : l'étanchéité dépend des joints, qui vieillissent. Elle se vérifie lors des révisions ; la couronne doit être repoussée, ou vissée si elle est à vis, avant tout contact avec l'eau."),
    ("h2", "9. L'entretien"),
    ("p", "Les huiles du mouvement vieillissent. Une révision complète est généralement recommandée tous les cinq à sept ans, selon la marque et l'usage : démontage, nettoyage, lubrification, remplacement des joints et réglage."),
    ("h2", "10. Idées reçues fréquentes"),
    ("p", "« Une automatique se remonte toute seule, même posée. » Non : elle se remonte quand on la porte. Posée, elle s'arrête au bout de sa réserve de marche."),
    ("p", "« Une montre arrêtée est en panne. » Le plus souvent, son ressort est simplement désarmé : il suffit de la porter ou de la remonter."),
    ("p", "« Plus de rubis, meilleure montre. » Non : le nombre de rubis dépend de la construction et des fonctions."),
    ("p", "« Le magnétisme, c'est fini. » Non : les sources d'aimants se sont multipliées dans notre quotidien."),
    ("p", "Conseil pratique : sur de nombreux mouvements, il vaut mieux ne pas corriger la date quand la montre indique entre 21 h et 3 h, car le mécanisme de changement de date est alors en prise."),
]


def guide_docx(path: Path, image: Path):
    doc = Document()
    for kind, text in GUIDE:
        if kind == "h1":
            doc.add_heading(text, 0)
        elif kind == "h2":
            doc.add_heading(text, 1)
        elif kind == "img":
            doc.add_picture(str(image), width=Cm(15))
        else:
            doc.add_paragraph(text).paragraph_format.space_after = Pt(6)
    doc.save(path)


# ---------------------------------------------------------------- field notes (DOCX)

NOTES = [
    ("Combien de temps tient-elle si je ne la porte pas ?", "La durée de sa réserve de marche, souvent deux jours environ. Ensuite elle s'arrête : il suffit de la remonter et de la remettre à l'heure."),
    ("Faut-il la remonter tous les jours ?", "Une manuelle, oui, de préférence chaque matin à la même heure. Une automatique portée chaque jour se remonte seule."),
    ("Pourquoi l'aiguille des secondes glisse-t-elle ?", "Parce que le balancier bat plusieurs fois par seconde : l'aiguille avance par petits pas, 8 par seconde à 4 hertz."),
    ("Que veut dire « 28 800 alternances par heure » ?", "C'est la vitesse du balancier : 8 alternances par seconde, soit 4 hertz."),
    ("Que signifie « chronomètre » ?", "Que le mouvement a réussi des tests officiels de précision : entre −4 et +6 secondes par jour en moyenne."),
    ("Ma montre avance de 30 secondes par jour depuis que je la pose sur mon ordinateur portable.", "C'est très probablement du magnétisme. Une démagnétisation chez l'horloger règle le problème en quelques minutes."),
    ("Puis-je nager avec ?", "Seulement si son étanchéité le permet et si les joints ont été vérifiés récemment ; couronne repoussée ou vissée."),
    ("Elle s'est arrêtée dans le tiroir : elle est cassée ?", "Non, elle est simplement désarmée. On la remonte, puis on la remet à l'heure."),
]


def notes_docx(path: Path):
    doc = Document()
    doc.add_heading("Notes de terrain — questions de clients", 0)
    doc.add_paragraph("Questions entendues en boutique et réponses conseillées (contenu de démonstration inventé).")
    for q, a in NOTES:
        doc.add_heading(q, 2)
        doc.add_paragraph(a)
    doc.save(path)


# ---------------------------------------------------------------- sales deck (PPTX)

SLIDES = [
    ("Expliquer le mouvement mécanique à un client", "Argumentaire conseillers — contenu de démonstration", None),
    ("Le principe en une phrase", "Une montre mécanique est un petit moteur à ressort, régulé par un balancier qui bat plusieurs fois par seconde.\nPas de pile : seulement de la mécanique.", None),
    ("Les cinq fonctions", "1 Énergie — 2 Transmission — 3 Distribution — 4 Régulation — 5 Affichage", "image"),
    ("La méthode en trois temps", "Écouter : partir de la question du client.\nMontrer : retourner la montre, faire observer le balancier qui oscille.\nRelier : du mécanisme au bénéfice (savoir-faire, durabilité, transmission).", None),
    ("Faire observer", "Fond transparent : le balancier qui bat, la masse oscillante qui tourne.\nLaisser le client tenir la montre et écouter le tic-tac.", None),
    ("Objection : « moins précis qu'un quartz »", "Reconnaître : oui, le quartz est plus précis.\nExpliquer : la mécanique, c'est un savoir-faire, sans pile, réparable pendant des décennies.\nRecadrer : quelques secondes par jour suffisent largement au quotidien.", None),
    ("Objection : « elle s'est arrêtée »", "Ce n'est pas une panne : la réserve de marche est épuisée.\nPorter la montre ou la remonter, puis la remettre à l'heure.", None),
    ("Objection : « c'est fragile »", "Un antichoc protège le balancier des chocs du quotidien.\nÉtanchéité vérifiée à chaque révision.\nÉviter les aimants : haut-parleurs, fermoirs magnétiques.", None),
    ("À la remise de la montre", "Remonter avant le premier port.\nNe pas corriger la date entre 21 h et 3 h.\nRévision tous les cinq à sept ans, selon la marque.", None),
    ("À retenir", "Cinq fonctions, une phrase simple, une démonstration visuelle.\nToujours partir de la question du client.", None),
]


def deck_pptx(path: Path, image: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = PCm(33.867), PCm(19.05)
    for i, (title, body, extra) in enumerate(SLIDES):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        for p in slide.placeholders[1].text_frame.paragraphs:
            for r in p.runs:
                r.font.size = PPt(22)
        if extra == "image":
            slide.shapes.add_picture(str(image), PCm(9), PCm(8.5), width=PCm(16))
        slide.notes_slide.notes_text_frame.text = f"Diapositive {i + 1} — notes de l'animateur."
    prs.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        image = tmpdir / "schema-mouvement.png"
        schematic(image)
        guide = tmpdir / "guide-mouvement.docx"
        guide_docx(guide, image)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(guide)],
                       check=True, capture_output=True)
        deck_pptx(OUT / "argumentaire.pptx", image)
        notes_docx(OUT / "notes-terrain.docx")
    print(f"sources written to {OUT}")


if __name__ == "__main__":
    main()
